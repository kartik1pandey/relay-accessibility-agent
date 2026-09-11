"""
build_pptx.py — generates the editable Unstop "Short Pitch Presentation
in PPT" upload from slides_content.py. Run: python build_pptx.py
"""

import os

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

from slides_content import SLIDES

BG = RGBColor(0xF5, 0xF7, 0xF4)
INK = RGBColor(0x10, 0x15, 0x1A)
INK_SOFT = RGBColor(0x3F, 0x4A, 0x47)
ACCENT = RGBColor(0x1B, 0x6B, 0x6E)
ACCENT_TEXT = RGBColor(0x0B, 0x45, 0x47)
LINE = RGBColor(0xD8, 0xDE, 0xD8)
DARK_BG = RGBColor(0x10, 0x15, 0x1A)
DARK_FG = RGBColor(0xF5, 0xF7, 0xF4)
DARK_MUTED = RGBColor(0x9F, 0xB3, 0xAC)

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)  # 16:9


def _blank_slide(prs, bg=BG):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # fully blank layout
    rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    rect.fill.solid()
    rect.fill.fore_color.rgb = bg
    rect.line.fill.background()
    rect.shadow.inherit = False
    # send background rect behind everything else added later
    spTree = slide.shapes._spTree
    spTree.remove(rect._element)
    spTree.insert(2, rect._element)
    return slide


def _textbox(slide, left, top, width, height):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    return box, tf


def _set_run(run, text, size, color, bold=False, font="Segoe UI"):
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = font


def build_title_slide(prs, s):
    slide = _blank_slide(prs, DARK_BG)
    box, tf = _textbox(slide, Inches(1), Inches(2.6), Inches(11.3), Inches(1.6))
    p = tf.paragraphs[0]
    r = p.add_run()
    _set_run(r, s["title"], 80, DARK_FG, bold=True)

    box2, tf2 = _textbox(slide, Inches(1), Inches(3.9), Inches(11.3), Inches(1.0))
    p2 = tf2.paragraphs[0]
    r2 = p2.add_run()
    _set_run(r2, s["subtitle"], 26, DARK_MUTED)

    box3, tf3 = _textbox(slide, Inches(1), Inches(6.6), Inches(11.3), Inches(0.6))
    p3 = tf3.paragraphs[0]
    r3 = p3.add_run()
    _set_run(r3, s["footer"], 18, ACCENT, bold=True)


def _eyebrow_and_title(slide, eyebrow, title):
    box, tf = _textbox(slide, Inches(0.8), Inches(0.55), Inches(11.7), Inches(0.5))
    r = tf.paragraphs[0].add_run()
    _set_run(r, eyebrow, 16, ACCENT_TEXT, bold=True)

    line = slide.shapes.add_connector(1, Inches(0.8), Inches(1.05), Inches(12.5), Inches(1.05))
    line.line.color.rgb = LINE
    line.line.width = Pt(1)

    box2, tf2 = _textbox(slide, Inches(0.8), Inches(1.2), Inches(11.7), Inches(1.3))
    r2 = tf2.paragraphs[0].add_run()
    _set_run(r2, title, 32, INK, bold=True)


def build_content_slide(prs, s):
    slide = _blank_slide(prs)
    _eyebrow_and_title(slide, s["eyebrow"], s["title"])
    box, tf = _textbox(slide, Inches(0.9), Inches(2.6), Inches(11.3), Inches(4.4))
    for i, bullet in enumerate(s["bullets"]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(18)
        r = p.add_run()
        _set_run(r, f"—  {bullet}", 20, INK_SOFT)


def build_diagram_slide(prs, s):
    slide = _blank_slide(prs)
    _eyebrow_and_title(slide, s["eyebrow"], s["title"])
    steps = s["steps"]
    top = Inches(2.3)
    step_h = Inches(0.62)
    gap = Inches(0.05)
    for i, step in enumerate(steps):
        y = Emu(int(top) + i * int(step_h + gap))
        is_gate = "Guardrail" in step
        box, tf = _textbox(slide, Inches(2.0), y, Inches(9.3), step_h)
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        r = tf.paragraphs[0].add_run()
        _set_run(r, step, 20 if is_gate else 18, ACCENT_TEXT if is_gate else INK, bold=is_gate)


def build_table_slide(prs, s):
    slide = _blank_slide(prs)
    _eyebrow_and_title(slide, s["eyebrow"], s["title"])
    rows, cols = len(s["rows"]) + 1, len(s["columns"])
    left, top = Inches(0.9), Inches(2.3)
    width, height = Inches(11.3), Inches(4.0)
    gtable = slide.shapes.add_table(rows, cols, left, top, width, height)
    table = gtable.table
    table.columns[0].width = Inches(2.0)
    table.columns[1].width = Inches(6.3)
    table.columns[2].width = Inches(3.0)

    for c, col_name in enumerate(s["columns"]):
        cell = table.cell(0, c)
        cell.text = col_name
        cell.fill.solid()
        cell.fill.fore_color.rgb = INK
        for p in cell.text_frame.paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.size = Pt(15)
                r.font.color.rgb = BG

    for ridx, row in enumerate(s["rows"], start=1):
        is_relay = row[0] == "Relay"
        for c, val in enumerate(row):
            cell = table.cell(ridx, c)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(0xEA, 0xEF, 0xE9) if is_relay else BG
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(13)
                for r in p.runs:
                    r.font.size = Pt(13)
                    r.font.bold = is_relay
                    r.font.color.rgb = ACCENT_TEXT if is_relay else INK_SOFT

    if s.get("note"):
        box, tf = _textbox(slide, Inches(0.9), Inches(6.5), Inches(11.3), Inches(0.6))
        r = tf.paragraphs[0].add_run()
        _set_run(r, s["note"], 13, INK_SOFT)


def build_stats_slide(prs, s):
    slide = _blank_slide(prs)
    _eyebrow_and_title(slide, s["eyebrow"], s["title"])
    n = len(s["stats"])
    col_w = Inches(11.3) / n
    for i, stat in enumerate(s["stats"]):
        left = Inches(0.9) + col_w * i
        box, tf = _textbox(slide, left, Inches(2.8), col_w, Inches(1.4))
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        r = tf.paragraphs[0].add_run()
        _set_run(r, stat["value"], 60, ACCENT_TEXT, bold=True)

        box2, tf2 = _textbox(slide, left, Inches(4.15), col_w, Inches(0.5))
        tf2.paragraphs[0].alignment = PP_ALIGN.CENTER
        r2 = tf2.paragraphs[0].add_run()
        _set_run(r2, stat["label"], 18, INK_SOFT)

    box3, tf3 = _textbox(slide, Inches(1.3), Inches(5.3), Inches(10.7), Inches(1.6))
    tf3.paragraphs[0].alignment = PP_ALIGN.CENTER
    r3 = tf3.paragraphs[0].add_run()
    _set_run(r3, s["note"], 16, INK_SOFT)


def build_closing_slide(prs, s):
    slide = _blank_slide(prs, DARK_BG)
    box, tf = _textbox(slide, Inches(1), Inches(2.9), Inches(11.3), Inches(1.2))
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    r = tf.paragraphs[0].add_run()
    _set_run(r, s["title"], 60, DARK_FG, bold=True)

    for i, line in enumerate(s["lines"]):
        box2, tf2 = _textbox(slide, Inches(1), Inches(4.3 + i * 0.5), Inches(11.3), Inches(0.5))
        tf2.paragraphs[0].alignment = PP_ALIGN.CENTER
        r2 = tf2.paragraphs[0].add_run()
        _set_run(r2, line, 20, DARK_MUTED)


BUILDERS = {
    "title": build_title_slide,
    "content": build_content_slide,
    "diagram": build_diagram_slide,
    "table": build_table_slide,
    "stats": build_stats_slide,
    "closing": build_closing_slide,
}


def main():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    for s in SLIDES:
        BUILDERS[s["kind"]](prs, s)
    out_path = os.path.join(os.path.dirname(__file__), "Relay_Pitch_Deck.pptx")
    prs.save(out_path)
    print(f"wrote {out_path}")


if __name__ == "__main__":
    main()
