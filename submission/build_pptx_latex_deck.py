r"""
build_pptx_latex_deck.py — generates an editable .pptx matching
Relay_Pitch_Deck_LaTeX.tex slide-for-slide (same 14 slides, same section
grouping, same content), styled to echo that Beamer theme: white
background, a single blue-violet "structure" color applied to all text
(titles, bullets, table content — matching the theme's
`\setbeamercolor{normal text}{fg=structure.fg}`), a thin rule under each
frame title, and circular bullet markers.

This is a separate, hand-written content mirror of the .tex file, not a
shared-data-driven generator — the LaTeX deck's prose (bold spans,
inline math, tikz diagrams) doesn't fit the older slides_content.py
schema, so keeping the two in sync means updating both when the content
changes. If Relay_Pitch_Deck_LaTeX.tex changes, update this file to match.

Run: python build_pptx_latex_deck.py
"""

import os

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

ACCENT = RGBColor(0x2D, 0x2D, 0x96)   # the theme's "structure" blue-violet
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0x2D, 0x2D, 0x96)
BOXFILL = RGBColor(0xFF, 0xFF, 0xFF)

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)  # 16:9
FONT = "Segoe UI"

OUT_PATH = os.path.join(os.path.dirname(__file__), "Relay_Pitch_Deck_LaTeX.pptx")


def _blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _textbox(slide, left, top, width, height):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    return box, tf


def _run(p, text, size=18, bold=False, italic=False, color=ACCENT, font=FONT):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = font
    return r


def _rich_paragraph(tf, segments, size=18, first=False, bullet=False, level=0):
    """segments: list of (text, bold, italic) tuples forming one paragraph."""
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_after = Pt(10)
    p.level = level
    if bullet:
        # python-pptx has no direct bullet API on a blank textbox; emulate
        # with a leading glyph, matching the theme's circular markers.
        _run(p, ("      " * level) + ("○  " if level == 0 else "-  "), size=size)
    for text, bold, italic in segments:
        _run(p, text, size=size, bold=bold, italic=italic)
    return p


def _frame_title(slide, title):
    box, tf = _textbox(slide, Inches(0.6), Inches(0.35), Inches(12.1), Inches(0.6))
    p = tf.paragraphs[0]
    _run(p, title, size=24, bold=True)
    line = slide.shapes.add_connector(1, Inches(0.6), Inches(0.95), Inches(12.73), Inches(0.95))
    line.line.color.rgb = LINE
    line.line.width = Pt(1)


def _box(slide, left, top, width, height, text, size=11):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = BOXFILL
    shape.line.color.rgb = ACCENT
    shape.line.width = Pt(1)
    shape.shadow.inherit = False
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    for i, line in enumerate(text.split("\n")):
        pp = p if i == 0 else tf.add_paragraph()
        pp.alignment = PP_ALIGN.CENTER
        _run(pp, line, size=size)
    return shape


def _arrow(slide, x1, y1, x2, y2):
    conn = slide.shapes.add_connector(1, x1, y1, x2, y2)
    conn.line.color.rgb = ACCENT
    conn.line.width = Pt(1.5)
    # python-pptx has no public arrowhead API; add the tail-end triangle
    # marker directly via the underlying XML (a documented workaround,
    # not exposed on the Line object itself).
    ln = conn.line._get_or_add_ln()
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle"})
    ln.append(tail)


# ---------------------------------------------------------------- slides

def slide_title(prs):
    slide = _blank_slide(prs)
    box, tf = _textbox(slide, Inches(1), Inches(2.3), Inches(11.3), Inches(1.0))
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    _run(p, "Relay", size=54, bold=True)

    box2, tf2 = _textbox(slide, Inches(1), Inches(3.2), Inches(11.3), Inches(1.1))
    p2 = tf2.paragraphs[0]
    p2.alignment = PP_ALIGN.CENTER
    _run(p2, "An On-Device, Voice-Controlled, Screen-Reading Accessibility Agent", size=26, bold=True)

    for i, (text, size) in enumerate([
        ("Kartik Pandey", 18),
        ("IIIT Naya Raipur", 13),
        ("Snapdragon AI Lab Build & Present Challenge, September 2026", 16),
    ]):
        box3, tf3 = _textbox(slide, Inches(1), Inches(4.6 + i * 0.5), Inches(11.3), Inches(0.5))
        p3 = tf3.paragraphs[0]
        p3.alignment = PP_ALIGN.CENTER
        _run(p3, text, size=size)


def slide_outline(prs):
    slide = _blank_slide(prs)
    _frame_title(slide, "Outline")
    box, tf = _textbox(slide, Inches(1.2), Inches(1.5), Inches(10.9), Inches(5.2))
    sections = [
        "Introduction", "Pipeline Overview", "Screen Understanding",
        "The Guardrail Engine", "Implementation", "Results",
        "Limitations & Future Work",
    ]
    for i, s in enumerate(sections):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(16)
        _run(p, s, size=22)


def slide_bullets(prs, title, bullets, block=None):
    """bullets: list of (segments, level) where segments is a list of (text,bold,italic)."""
    slide = _blank_slide(prs)
    _frame_title(slide, title)
    box, tf = _textbox(slide, Inches(0.9), Inches(1.4), Inches(11.5), Inches(4.6))
    first = True
    for segments, level in bullets:
        _rich_paragraph(tf, segments, size=18 if level == 0 else 16, first=first, bullet=True, level=level)
        first = False

    if block:
        block_title, block_text = block
        # Real bug found by actually rendering this slide: a fixed
        # per-bullet height estimate undershot how much a bullet
        # genuinely wraps to at this width/font size, so the block
        # overlapped the last line of bullet text. Estimating by
        # character count (roughly how many lines each bullet's text
        # will actually wrap to at ~95 chars/line) is a closer match to
        # what word_wrap actually produces.
        est_lines = sum(max(1, -(-sum(len(t) for t, _, _ in segs) // 95)) for segs, _ in bullets)
        top = Inches(1.4) + Inches(0.42) * est_lines + Inches(0.5)
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), top, Inches(11.5), Inches(1.3))
        shape.fill.solid()
        shape.fill.fore_color.rgb = WHITE
        shape.line.color.rgb = ACCENT
        shape.line.width = Pt(1)
        shape.shadow.inherit = False
        btf = shape.text_frame
        btf.word_wrap = True
        btf.margin_left = Inches(0.15)
        btf.margin_right = Inches(0.15)
        bp = btf.paragraphs[0]
        _run(bp, block_title, size=15, bold=True)
        bp2 = btf.add_paragraph()
        _run(bp2, block_text, size=14)


def slide_pipeline(prs):
    slide = _blank_slide(prs)
    _frame_title(slide, "End-to-End Pipeline")

    w, h, gap = Inches(2.1), Inches(0.85), Inches(0.35)
    top1 = Inches(1.5)
    top2 = top1 + h + Inches(0.8)
    xs = [Inches(0.7) + i * (w + gap) for i in range(4)]

    row1 = ["Voice\ncommand", "Whisper\nspeech-to-text", "Screen\nunderstanding", "Reasoning\nmodel"]
    row2 = ["MeloTTS\nspeaks result", "Action\nexecutes", "Guardrail\ncheck", ""]

    boxes1, boxes2 = [], []
    for i, label in enumerate(row1):
        boxes1.append(_box(slide, xs[i], top1, w, h, label))
    for i, label in enumerate(row2):
        if label:
            boxes2.append((xs[i], _box(slide, xs[i], top2, w, h, label)))
        else:
            boxes2.append((xs[i], None))

    # row1 left-to-right arrows
    for i in range(3):
        _arrow(slide, xs[i] + w, top1 + h // 2, xs[i + 1], top1 + h // 2)
    # down from reasoning (col 3) to guardrail (col 2, row2)
    _arrow(slide, xs[3] + w // 2, top1 + h, xs[2] + w // 2, top2)
    # row2 right-to-left: guardrail(col2) -> action(col1) -> tts(col0)
    _arrow(slide, xs[2], top2 + h // 2, xs[1] + w, top2 + h // 2)
    _arrow(slide, xs[1], top2 + h // 2, xs[0] + w, top2 + h // 2)

    box, tf = _textbox(slide, Inches(0.9), top2 + h + Inches(0.5), Inches(11.5), Inches(2.0))
    _rich_paragraph(tf, [(
        "Screen understanding checks the UI Automation accessibility tree first; "
        "only falls back to a vision-language model when that tree doesn't have "
        "enough named, enabled controls to work with.", False, False)], first=True, bullet=True, size=16)
    _rich_paragraph(tf, [(
        "The guardrail check never changes: two-tier ALLOW / CONFIRM / BLOCK "
        "against pattern-matched control names, evaluated before any action executes.",
        False, False)], bullet=True, size=16)


def slide_decision_flow(prs):
    slide = _blank_slide(prs)
    _frame_title(slide, "The Full Decision Flow")

    guard = _box(slide, Inches(1.0), Inches(1.6), Inches(3.6), Inches(1.0), "Guardrail check\n(pattern match on name)")
    confirm = _box(slide, Inches(1.0), Inches(3.6), Inches(3.6), Inches(1.0), "Spoken confirmation\n(“yes” required)")

    box, tf = _textbox(slide, Inches(5.3), Inches(1.5), Inches(4.5), Inches(0.7))
    _run(tf.paragraphs[0], "ALLOW → act immediately", size=13)
    box2, tf2 = _textbox(slide, Inches(5.3), Inches(2.35), Inches(4.5), Inches(0.7))
    _run(tf2.paragraphs[0], "BLOCK → never executes", size=13)
    box3, tf3 = _textbox(slide, Inches(5.3), Inches(3.75), Inches(4.5), Inches(1.0))
    p3 = tf3.paragraphs[0]
    _run(p3, "“yes” → act", size=13)
    p3b = tf3.add_paragraph()
    _run(p3b, "silence/“no” → cancel", size=13)

    _arrow(slide, Inches(4.6), Inches(1.9), Inches(5.3), Inches(1.75))
    _arrow(slide, Inches(4.6), Inches(2.0), Inches(5.3), Inches(2.6))
    _arrow(slide, Inches(2.8), Inches(2.6), Inches(2.8), Inches(3.6))
    _arrow(slide, Inches(4.6), Inches(4.1), Inches(5.3), Inches(4.15))

    box4, tf4 = _textbox(slide, Inches(0.9), Inches(5.1), Inches(11.5), Inches(1.5))
    _rich_paragraph(tf4, [("Most decisions resolve immediately (ALLOW or BLOCK).", False, False)], first=True, bullet=True, size=16)
    _rich_paragraph(tf4, [("Only genuinely consequential actions pay the extra cost of a spoken confirmation round-trip.", False, False)], bullet=True, size=16)


def slide_table(prs, title, columns, rows, caption, col_widths, extra_bullets=None):
    slide = _blank_slide(prs)
    _frame_title(slide, title)

    left, top = Inches(0.9), Inches(1.5)
    n_rows, n_cols = len(rows) + 1, len(columns)
    total_w = sum(col_widths)
    tbl_shape = slide.shapes.add_table(n_rows, n_cols, left, top, Inches(total_w), Inches(0.5 * n_rows))
    table = tbl_shape.table
    for i, w in enumerate(col_widths):
        table.columns[i].width = Inches(w)

    for c, name in enumerate(columns):
        cell = table.cell(0, c)
        cell.text = name
        cell.fill.solid()
        cell.fill.fore_color.rgb = WHITE
        for p in cell.text_frame.paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.size = Pt(13)
                r.font.color.rgb = ACCENT
                r.font.name = FONT

    for ridx, row in enumerate(rows, start=1):
        is_relay = row[0] == "Relay"
        for c, val in enumerate(row):
            cell = table.cell(ridx, c)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(12)
                    r.font.bold = is_relay
                    r.font.color.rgb = ACCENT
                    r.font.name = FONT

    cap_top = top + Inches(0.5 * n_rows) + Inches(0.2)
    box, tf = _textbox(slide, left, cap_top, Inches(total_w), Inches(0.6))
    p = tf.paragraphs[0]
    _run(p, caption, size=12, italic=True)

    if extra_bullets:
        box2, tf2 = _textbox(slide, Inches(0.9), cap_top + Inches(0.8), Inches(11.5), Inches(2.0))
        first = True
        for text in extra_bullets:
            _rich_paragraph(tf2, [(text, False, False)], first=first, bullet=True, size=16)
            first = False


def slide_closing(prs):
    slide = _blank_slide(prs)
    _frame_title(slide, "Honest Roadmap")
    box, tf = _textbox(slide, Inches(0.9), Inches(1.5), Inches(11.5), Inches(4.0))
    bullets = [
        [("Open: ", True, False), ("Outlook's Chromium accessibility gap — the same limitation as Edge, but it's a desktop app the DevTools Protocol can't attach to. Unsolved, documented as such.", False, False)],
        [("Open: ", True, False), ("full GenieX-runtime, on-device latency for the reasoning model — needs physical Snapdragon hardware, not yet measured.", False, False)],
        [("Done: ", True, False), ("18 adversarial guardrail tests, a fixed task-list harness (100% pass rate, 5/5), and short-term reasoning memory (last 3 turns).", False, False)],
        [("This submission was built without physical Snapdragon hardware — every claim in this deck is labeled by exactly how it was verified.", False, False)],
    ]
    first = True
    for segs in bullets:
        _rich_paragraph(tf, segs, size=17, first=first, bullet=True)
        first = False

    box2, tf2 = _textbox(slide, Inches(0.9), Inches(6.2), Inches(11.5), Inches(0.5))
    p2 = tf2.paragraphs[0]
    p2.alignment = PP_ALIGN.CENTER
    _run(p2, "github.com/kartik1pandey/relay-accessibility-agent", size=15, italic=True)


def main():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    slide_title(prs)
    slide_outline(prs)

    slide_bullets(prs, "The Problem", [
        ([("People with ", False, False), ("limited hand mobility or low vision", True, False), (" can't reliably operate a Windows PC through a mouse and a screen alone.", False, False)], 0),
        ([("Existing screen readers ", False, False), ("describe", False, True), (" a screen but don't act on it; existing automation agents act, but weren't built around this population's safety needs.", False, False)], 0),
        ([("Idea: ", True, False), ("a voice-first agent that both understands the screen and acts on it — with a person narrating every step, and a spoken confirmation gate on anything consequential.", False, False)], 0),
    ], block=("Design goals", "On-device only, no cloud round-trip for speech or screen understanding · never act on anything sensitive without a spoken “yes” · scoped assistive technology, not a general autonomous agent"))

    slide_bullets(prs, "Why Not Just Use a General Computer-Use Agent?", [
        ([("General autonomous agents can already read a screen and click or type on it.", False, False)], 0),
        ([("But the cost of being wrong is asymmetric:", False, False)], 0),
        ([("Agent mis-clicks a harmless button → annoying, but recoverable.", False, False)], 1),
        ([("Agent confidently taps “Send” / “Delete” / “Pay” on a misread screen → real, sometimes irreversible harm.", False, False)], 1),
        ([("⇒ the whole design is biased toward ", False, False), ("safety over autonomy", True, False), (": nothing consequential executes without a spoken confirmation, ever.", False, False)], 0),
    ])

    slide_pipeline(prs)

    slide_bullets(prs, "The Fast Path — UI Automation Tree", [
        ([("Reads the live accessibility tree first: named, enabled controls, no model call needed.", False, False)], 0),
        ([("Verified against real, live windows on this machine:", False, False)], 0),
        ([("Media Player", True, False), (" — 211 elements, 113 named + enabled, real actionable content (Search, Open file(s), Recent media).", False, False)], 1),
        ([("Notepad", True, False), (" — a simple-app proxy, 24 elements, a real scroll action executed successfully.", False, False)], 1),
        ([("Falls back to a vision-language model only when the tree genuinely doesn't have enough to work with — not the common case, the exception.", False, False)], 0),
    ])

    slide_bullets(prs, "The Chromium Gap — Found, Verified, and Fixed", [
        ([("Edge and Outlook", True, False), (" are both Chromium/WebView2-based: walking their UI Automation tree 12–14 levels deep surfaces window chrome only — ", False, False), ("zero", True, False), (" real page content, confirmed, not assumed.", False, False)], 0),
        ([("Fixed for Edge", True, False), (" via the Chrome DevTools Protocol: real page content read from a live Wikipedia page, a real value typed into the real search box, and the write independently confirmed by re-reading the live DOM afterward.", False, False)], 0),
        ([("Outlook's equivalent gap remains open — it's a desktop app, not a browser tab the DevTools Protocol can attach to.", False, False)], 0),
    ], block=("Every claim here is labeled by how it was verified", "Run on real hardware/software, confirmed from real source code, or explicitly marked as not yet possible without Snapdragon-class hardware."))

    slide_bullets(prs, "Two-Tier Decision", [
        ([("Every decision is checked against the target control's name ", False, False), ("before", False, True), (" it executes:", False, False)], 0),
        ([("BLOCK", True, False), ("      name matches a block pattern (irreversible)", False, False)], 1),
        ([("CONFIRM", True, False), ("   name matches a confirm pattern (consequential)", False, False)], 1),
        ([("ALLOW", True, False), ("      otherwise", False, False)], 1),
        ([("CONFIRM", True, False), (": send/submit/post/publish, pay/buy/transfer, delete/remove/erase, account or security changes — requires a spoken “yes” before proceeding.", False, False)], 0),
        ([("BLOCK", True, False), (": factory reset, reformat, disabling the firewall or antivirus — never executes, no override.", False, False)], 0),
        ([("Every decision is written to a persisted, append-only audit log.", False, False)], 0),
    ])

    slide_bullets(prs, "Adversarial Testing", [
        ([("18 adversarial test cases built from plausible real controls across all three target apps.", False, False)], 0),
        ([("Found a genuine gap: the pattern ", False, False), ("\\bdelete\\b", False, False), ("'s word boundary didn't match inside “deleted” — so a real control like “Empty deleted items folder” silently fell through to ALLOW.", False, False)], 0),
        ([("Fixed by widening ", False, False), ("delete/remove/erase", False, False), (" to also catch -ed/-ing/-ion forms — a widening, never a narrowing, by design: these patterns should get harder to change, not easier.", False, False)], 0),
    ])

    slide_decision_flow(prs)

    slide_bullets(prs, "What's Actually Real", [
        ([("Whisper speech-to-text", True, False), (" — real transcription of a known test clip, reproduced three independent ways.", False, False)], 0),
        ([("MeloTTS-EN", True, False), (" — real synthesis, real audio, played back through real speakers (five independent real blockers found and fixed to get there).", False, False)], 0),
        ([("The full agent.py orchestration loop", True, False), (" — a real mic recording → real transcription → real UI Automation tree read → real guardrail gate → real action → real speech, proven end to end against a live window. Only the NPU-dependent reasoning step is mocked.", False, False)], 0),
        ([("cdp_perception.py", True, False), (" — the Chrome DevTools Protocol fix for Edge, wired into agent.py's real orchestration and demonstrated live.", False, False)], 0),
    ])

    slide_table(
        prs, "Competitive Differentiation",
        ["Project", "Safety mechanism", "Enforcement"],
        [
            ["Deft", "None found anywhere in either backing repo", "N/A"],
            ["Orbit", "None — but never takes actions at all", "N/A"],
            ["ScreenSaathi", "Keyword list downgrades risky plans", "Code-level"],
            ["Sally", "A prompt asking the model nicely", "Model-enforced only"],
            ["Relay", "ALLOW / CONFIRM / BLOCK + spoken yes + audit log", "Code-level, pre-execution"],
        ],
        "Verified against each project's real, public source code — not a secondhand summary.",
        [2.3, 6.5, 2.5],
        extra_bullets=["Of five comparable projects, Relay is the only one with a deterministic, code-level safety gate that doesn't depend on a language model correctly following an instruction."],
    )

    slide_table(
        prs, "Real Snapdragon Hardware — Qualcomm AI Hub",
        ["Metric", "Value"],
        [
            ["Whisper-base encoder", "~49.1 ms"],
            ["Whisper-base decoder", "~3.7 ms"],
        ],
        "Compiled and profiled on a real Snapdragon X Elite CRD via Qualcomm AI Hub's hosted device farm.",
        [5.0, 3.0],
        extra_bullets=[
            "A real compile + profile job on actual target silicon — not a simulation, and an order of magnitude faster than this laptop's CPU.",
            "Precisely scoped: this is AI Hub's compile/profile path, not yet a full GenieX runtime measurement, which needs physical device access.",
        ],
    )

    slide_closing(prs)

    prs.save(OUT_PATH)
    print(f"wrote {OUT_PATH} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
