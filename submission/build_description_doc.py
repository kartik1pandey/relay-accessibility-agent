"""
build_description_doc.py — generates the "Brief Project Description"
upload (docx, per Unstop's accepted pdf/doc/docx types) for the Snapdragon
AI Lab Build & Present Challenge submission. Content is a condensed,
consistent version of WRITEUP.md — every claim traceable to CLAUDE.md's
verified findings.

Run: python build_description_doc.py
"""

import os

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

INK = RGBColor(0x10, 0x15, 0x1A)
INK_SOFT = RGBColor(0x3F, 0x4A, 0x47)
ACCENT_TEXT = RGBColor(0x0B, 0x45, 0x47)

OUT_PATH = os.path.join(os.path.dirname(__file__), "Relay_Project_Description.docx")


def _style_run(run, size=11, color=INK, bold=False, italic=False):
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = "Segoe UI"


def _heading(doc, text):
    p = doc.add_paragraph()
    p.space_before = Pt(14)
    p.space_after = Pt(4)
    r = p.add_run(text)
    _style_run(r, size=14, color=ACCENT_TEXT, bold=True)


def _body(doc, text):
    p = doc.add_paragraph()
    p.space_after = Pt(8)
    r = p.add_run(text)
    _style_run(r, size=11, color=INK_SOFT)


def _bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.space_after = Pt(4)
    r = p.add_run(text)
    _style_run(r, size=11, color=INK_SOFT)


def main():
    doc = Document()

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = title_p.add_run("Relay")
    _style_run(r, size=28, color=INK, bold=True)

    sub_p = doc.add_paragraph()
    r = sub_p.add_run("An on-device, voice-controlled, screen-reading accessibility agent")
    _style_run(r, size=13, color=INK_SOFT, italic=True)

    event_p = doc.add_paragraph()
    r = event_p.add_run("Snapdragon AI Lab Build & Present Challenge")
    _style_run(r, size=11, color=ACCENT_TEXT, bold=True)

    _heading(doc, "The problem")
    _body(
        doc,
        "People with limited hand mobility or low vision can't reliably operate a "
        "Windows PC through a mouse and a screen alone. Existing screen readers "
        "describe a screen but don't act on it; existing automation agents act, "
        "but weren't designed around this population's safety needs. Relay is a "
        "voice-controlled, screen-reading agent, running entirely on-device, that "
        "narrates its own actions and gates anything consequential behind a "
        "spoken confirmation. Someone says “open my email and read me the new "
        "ones” — that's the whole interaction.",
    )
    _body(
        doc,
        "This is deliberately not a general autonomous computer-use agent. It's "
        "scoped assistive technology: a person narrates every step, and nothing "
        "consequential happens without a spoken “yes.” That framing is a "
        "feature of the design, not a limitation.",
    )

    _heading(doc, "How it works")
    _body(
        doc,
        "Voice command → Whisper (speech-to-text) → screen understanding "
        "(the UI Automation accessibility tree first; a vision-language model "
        "only when that tree doesn't have enough named, enabled controls to "
        "work with) → a small local reasoning model decides the next action "
        "→ a two-tier guardrail check (ALLOW / CONFIRM with a spoken yes / "
        "BLOCK) → the action executes against the real UI element → MeloTTS "
        "speaks the result. Every guardrail decision is written to a persisted, "
        "append-only audit log.",
    )

    _heading(doc, "What makes it different")
    _body(
        doc,
        "Four comparable, real accessibility-agent projects were researched directly "
        "from their public source code (not a secondhand summary): of the five "
        "projects compared, Relay is the only one with a deterministic, code-level "
        "safety gate that doesn't depend on a language model correctly following an "
        "instruction. One competitor's entire safety mechanism is a prompt asking its "
        "model nicely not to tap Send or Delete — with no code-level backstop at all.",
    )
    _bullet(doc, "Two-tier ALLOW / CONFIRM / BLOCK guardrail, pattern-matched pre-execution, with a spoken-yes gate and a persisted audit log.")
    _bullet(doc, "Two distinct, WCAG AAA-compliant visual surfaces: a 7-state end-user runtime UI and a guardrail-admin UI that reads the real rules from code, not a hand-maintained copy.")
    _bullet(doc, "A real, working fix for a genuine architectural gap: Chromium-based apps (Edge, Outlook) don't expose real page content via Windows UI Automation — found, verified, and fixed for Edge using the Chrome DevTools Protocol.")

    _heading(doc, "What's actually been verified, not just claimed")
    _bullet(doc, "Whisper speech-to-text: real transcription of a known test clip, reproduced three independent ways.")
    _bullet(doc, "MeloTTS-EN text-to-speech: real synthesis, real audio, played back through real speakers.")
    _bullet(doc, "The full agent.py orchestration loop, end to end against live hardware: a real microphone recording → real transcription → a real live UI Automation tree read → the real guardrail gate → a real action → real speech — with only the NPU-dependent reasoning step mocked.")
    _bullet(doc, "The Chrome DevTools Protocol fix for Edge, demonstrated live: a real Wikipedia search box found, written to, and the write independently confirmed by re-reading the live page.")
    _bullet(doc, "18 adversarial guardrail test cases across all three target apps, and a fixed task-list harness passing 100% (5/5) against live windows.")
    _bullet(doc, "Real Snapdragon X Elite hardware numbers via Qualcomm AI Hub's hosted device farm: Whisper-base compiled and profiled at ~49.1ms (encoder) / ~3.7ms (decoder) per inference — a real compile-and-profile job on actual target silicon, precisely distinguished from a full GenieX runtime measurement, which needs physical device access this build doesn't have yet.")

    _heading(doc, "Honest scope and open questions")
    _bullet(doc, "Outlook's equivalent Chromium accessibility gap remains open — it's a desktop app, not a browser tab the DevTools Protocol can attach to.")
    _bullet(doc, "Full on-device GenieX-runtime latency for the reasoning model is not yet measured — that requires physical Snapdragon hardware, which this submission was built without.")
    _bullet(doc, "This project deliberately targets three genuinely different interaction shapes (mail, browser, media player) rather than expanding to more apps for its own sake.")

    _heading(doc, "Repository")
    _body(doc, "github.com/kartik1pandey/relay-accessibility-agent")

    doc.save(OUT_PATH)
    print(f"wrote {OUT_PATH}")


if __name__ == "__main__":
    main()
